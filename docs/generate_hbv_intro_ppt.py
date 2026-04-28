from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT_PATH = Path(__file__).resolve().parent / "HBV_model_intro_technical.pptx"


def set_font(run, size, bold=False, color=(0, 0, 0), name="Microsoft YaHei"):
    font = run.font
    font.name = name
    font.size = Pt(size)
    font.bold = bold
    font.color.rgb = RGBColor(*color)


def add_textbox(slide, left, top, width, height, text, size=14, bold=False,
                color=(0, 0, 0), fill=None, line=None, align=PP_ALIGN.LEFT,
                radius_shape=None, margin=0.08, name="Microsoft YaHei"):
    shape_type = radius_shape or MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE
    box = slide.shapes.add_shape(shape_type, left, top, width, height)

    if fill is None:
        box.fill.background()
    else:
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(*fill)

    if line is None:
        box.line.fill.background()
    else:
        box.line.color.rgb = RGBColor(*line)
        box.line.width = Pt(1.2)

    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_font(run, size=size, bold=bold, color=color, name=name)
    return box


def add_bullets(slide, left, top, width, height, title, bullets,
                title_size=16, bullet_size=11.5, fill=(248, 250, 252),
                line=(203, 213, 225)):
    box = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(*fill)
    box.line.color.rgb = RGBColor(*line)
    box.line.width = Pt(1.1)

    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    tf.margin_top = Inches(0.10)
    tf.margin_bottom = Inches(0.08)

    p0 = tf.paragraphs[0]
    p0.alignment = PP_ALIGN.LEFT
    r0 = p0.add_run()
    r0.text = title
    set_font(r0, size=title_size, bold=True, color=(15, 23, 42))

    for bullet in bullets:
        p = tf.add_paragraph()
        p.text = f"- {bullet}"
        p.level = 0
        p.alignment = PP_ALIGN.LEFT
        if p.runs:
            set_font(p.runs[0], size=bullet_size, color=(51, 65, 85))

    return box


def add_connector(slide, x1, y1, x2, y2, color=(148, 163, 184)):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.line.color.rgb = RGBColor(*color)
    line.line.width = Pt(2.2)
    line.line.end_arrowhead = True
    return line


def add_stage(slide, left, top, width, height, title, subtitle, fill):
    box = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(*fill)
    box.line.color.rgb = RGBColor(255, 255, 255)
    box.line.width = Pt(1.2)

    tf = box.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)

    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    r1 = p1.add_run()
    r1.text = title
    set_font(r1, size=17, bold=True, color=(255, 255, 255))

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run()
    r2.text = subtitle
    set_font(r2, size=10.5, color=(239, 246, 255))

    return box


def build_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])

    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(245, 247, 250)

    banner = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.9)
    )
    banner.fill.solid()
    banner.fill.fore_color.rgb = RGBColor(15, 23, 42)
    banner.line.fill.background()

    title_box = slide.shapes.add_textbox(Inches(0.45), Inches(0.18), Inches(8.9), Inches(0.42))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "基于 YAML 配置的水利模型通用对接框架"
    set_font(r, size=24, bold=True, color=(255, 255, 255))

    sub_box = slide.shapes.add_textbox(Inches(0.47), Inches(0.56), Inches(10.5), Inches(0.22))
    tf = sub_box.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "面向 HBV / SWAT / APEX 等模型的参数写入、仿真执行、结果提取与指标评估一体化流程"
    set_font(r, size=10.5, color=(203, 213, 225))

    add_textbox(
        slide, Inches(0.45), Inches(1.03), Inches(6.55), Inches(0.56),
        "定位：将外部水利模型封装为统一 Problem 接口，通过 YAML 描述参数、序列、指标与日志规则，直接接入 UQPyL 的率定/优化/不确定性分析。",
        size=12.5, bold=False, color=(30, 41, 59), fill=(226, 232, 240), line=(203, 213, 225),
        radius_shape=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, margin=0.10
    )

    add_bullets(
        slide, Inches(7.15), Inches(1.03), Inches(5.72), Inches(0.92),
        "配置层次",
        [
            "basic / parameters / functions / series / derived / objectives / constraints / diagnostics / reporter",
        ],
        title_size=14, bullet_size=10.2, fill=(255, 255, 255)
    )

    flow_top = Inches(2.12)
    box_w = Inches(1.93)
    box_h = Inches(0.92)
    gap = Inches(0.12)
    start_left = Inches(0.45)

    stages = [
        ("YAML 配置", "RunConfig + Pydantic 校验", (37, 99, 235)),
        ("参数映射", "ParamManager / WriteInHandler", (14, 165, 233)),
        ("模型运行", "tempRun + subprocess + timeout", (249, 115, 22)),
        ("结果提取", "ExtractSpec / ExprSpec / CallSpec", (147, 51, 234)),
        ("指标评估", "derived -> obj / con / diag", (220, 38, 38)),
        ("结果归档", "SQLite / CSV / error / warning", (22, 163, 74)),
    ]

    stage_boxes = []
    for idx, (title, subtitle, fill) in enumerate(stages):
        left = start_left + idx * (box_w + gap)
        stage_boxes.append(add_stage(slide, left, flow_top, box_w, box_h, title, subtitle, fill))

    for idx in range(len(stage_boxes) - 1):
        s1 = stage_boxes[idx]
        s2 = stage_boxes[idx + 1]
        x1 = s1.left + s1.width
        y1 = s1.top + s1.height // 2
        x2 = s2.left
        y2 = s2.top + s2.height // 2
        add_connector(slide, x1, y1, x2, y2)

    add_bullets(
        slide, Inches(0.45), Inches(3.32), Inches(4.18), Inches(2.85),
        "关键类与职责",
        [
            "`SimModel` 负责总控：创建运行队列、调度并行实例、回收目录、统一返回 `objs/cons`。",
            "`ParamManager` 负责 X->P 映射、参数缓存、文件匹配与写入任务注册。",
            "`WriteInHandler` 基于字节偏移覆盖定宽字段，支持相对/替换/增量三种模式。",
            "`SeriesExtractor` 将模型输出组织为 `sid_sim/sid_obs`，供后续 derived 与目标函数复用。",
            "`FunctionManager` 统一注册内置指标与外部 Python 函数，降低模型适配成本。",
        ],
        bullet_size=10.2
    )

    add_bullets(
        slide, Inches(4.78), Inches(3.32), Inches(4.2), Inches(2.85),
        "核心实现细节",
        [
            "`load_config` 使用 Pydantic 校验文件存在性、字段一致性以及依赖闭包是否完整。",
            "参数写入支持 exact / glob / regex 三种文件定位方式，适配复杂工程目录。",
            "序列提取支持 `colSpan` 定宽截取、`colNum` 分列提取，也支持表达式和函数调用。",
            "`Evaluator` 先计算 `derived`，再收集 objective / constraint / diagnostic 标量。",
            "失败时通过 `RunError(stage, code, target, message)` 统一描述错误来源与处理阶段。",
        ],
        bullet_size=10.2
    )

    add_bullets(
        slide, Inches(9.12), Inches(3.32), Inches(3.76), Inches(2.85),
        "并行与日志机制",
        [
            "`create_run_queue()` 会复制多个 `instance_i` 目录，避免并发写同一工程文件。",
            "`ThreadPoolExecutor` 并行提交样本，子进程超时后 kill 并写入 error 记录。",
            "`RunReporter` 后台线程按 batch/run 顺序落盘，输出 `results.db`、`summary.csv` 与序列 CSV。",
            "参数越界会写入 `warning.txt`，运行失败会写入 `error.txt`，便于排查问题样本。",
            "退出时执行 best-effort 清理，仅保留 backup 与结果文件。",
        ],
        bullet_size=10.0
    )

    footer = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(6.3), Inches(12.4), Inches(0.62)
    )
    footer.fill.solid()
    footer.fill.fore_color.rgb = RGBColor(219, 234, 254)
    footer.line.color.rgb = RGBColor(96, 165, 250)
    footer.line.width = Pt(1.2)

    tf = footer.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "一句话总结：该框架将外部水利模型封装为统一可调用的问题对象，可直接接入 UQPyL 的优化、率定与不确定性分析流程。"
    set_font(r, size=13, bold=True, color=(30, 64, 175))

    prs.save(str(OUT_PATH))


if __name__ == "__main__":
    build_ppt()
    print(OUT_PATH)
